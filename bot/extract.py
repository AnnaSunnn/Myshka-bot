import io

import docx
from pypdf import PdfReader


def extract_text(data: bytes, ext: str) -> str:
    if ext == "pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == "docx":
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="ignore")
